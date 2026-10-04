import os
import sys
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT / "_packaging_deps"))
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pandas as pd
import FINALChartSageV1 as module
from openpyxl import load_workbook
from paper_signal_replay import evaluate_position_exit_expression
from PySide6.QtTest import QTest


class _Signal:
    def connect(self, callback):
        self.callback = callback


class _FakePipeline:
    calls = []

    def __init__(self, **kwargs):
        self.kwargs = kwargs
        self.progress = _Signal()
        self.__class__.calls.append(kwargs)

    def _run_pipeline(self):
        name = (
            self.kwargs.get("finalset_name")
            or self.kwargs.get("exit_finalset_name")
            or next(iter(self.kwargs.get("member_names") or []), "")
            or next(iter(self.kwargs.get("exit_names") or []), "")
        )
        is_exit = bool(
            self.kwargs.get("exit_names")
            or self.kwargs.get("exit_finalset_name")
        )
        symbols = {
            "BUY A": ["AAA"],
            "BUY B": ["AAA", "BBB"],
            "BUY SET": ["CCC"],
            "EXIT X": ["BBB", "DDD"],
        }.get(name, [])
        signal = "EXIT / WEAKNESS" if is_exit else "BUY"
        return {
            "guidance_rows": [
                {
                    "Symbol": symbol,
                    "Signal": signal,
                    "Matched Members": name,
                    "Confirmation Count": 1,
                }
                for symbol in symbols
            ]
        }


class FinalStep3MultiSelectTests(unittest.TestCase):
    def setUp(self):
        app = module.QApplication.instance() or module.QApplication([])
        self.app = app
        rule = module.default_setup()
        self.templates = {
            "BUY A": dict(rule),
            "BUY B": dict(rule),
            "EXIT X": dict(rule),
            "BUY SET": {
                **module.default_setup(),
                "is_final_template": True,
                "finalset_side": "BUY",
                "finalset_op": "AND",
                "finalset_members": ["BUY A", "BUY B"],
            },
        }

    def test_loss_percent_parser_and_examples_use_positive_loss_wording(self):
        rule = module.RuleParser.parse("Loss % >= 5%")
        self.assertEqual(rule["left"]["name"], "Loss %")
        self.assertEqual(rule["operator"], ">=")
        self.assertEqual(rule["right"]["value"], 5.0)
        self.assertTrue(module.expression_has_pnl_rules(rule))
        examples = module.filter_examples_text()
        self.assertIn("Loss % >= 5%", examples)
        self.assertIn("Profit % <= -5%", examples)

    def test_position_rule_can_save_for_paper_without_zerodha_but_not_live_algo(self):
        loss_rule = module.RuleParser.parse("Loss % >= 5%")
        setup = module.default_setup()
        setup["universes"] = ["NIFTY 500"]
        setup["expression"] = module.group_node("AND", [loss_rule])
        dummy_window = SimpleNamespace(active_template_name=None)
        self.assertTrue(module.ScannerWindow._validate_template_setup(
            dummy_window, setup, silent=True
        ))
        self.assertEqual(setup["position_rule_scope"], "PAPER_BACKTEST_ONLY")
        self.assertEqual(setup["strategy_guide_role"], "EXIT")

        live_setup = dict(setup)
        live_setup["is_algo_template"] = True
        self.assertFalse(module.ScannerWindow._validate_template_setup(
            dummy_window, live_setup, silent=True
        ))

    def test_paper_history_combines_multiple_standalone_position_exits_as_or(self):
        templates = {
            name: module.default_setup()
            for name in ("FEEDER", "BUY A", "STOP LOSS", "TAKE PROFIT")
        }
        templates["FEEDER"]["universes"] = ["FIXED TEST SOURCE"]
        templates["STOP LOSS"]["expression"] = module.group_node("AND", [
            module.rule_node(
                module.field_operand("Loss %", "D"),
                ">=",
                module.value_operand(8.5),
            )
        ])
        templates["TAKE PROFIT"]["expression"] = module.group_node("AND", [
            module.rule_node(
                module.field_operand("Profit %", "D"),
                ">=",
                module.value_operand(16),
            )
        ])
        config = module.finalchartink_defaults()
        config.update({
            "recipe_name": "POSITION EXIT OR TEST",
            "feeder_templates": ["FEEDER"],
            "feeder_scope": "SAVED",
            "include_multirule": True,
            "buy_finalsets": ["BUY A"],
            "exit_finalsets": ["STOP LOSS", "TAKE PROFIT"],
            "as_of_mode": "historical",
            "as_of_date": "2025-10-03",
        })
        worker = module.FinalChartInkHistoryWorker(
            config,
            templates,
            {"FIXED TEST SOURCE": ("CHARTSAGE WATCHLIST", ["ABC"])},
            1,
        )

        def report(rows):
            return module.HistoryReport(
                rows_by_date=rows,
                daily_summary=[
                    {"date": "2025-10-01", "total": len(rows), "sectors": {}}
                ],
                total_symbols=1,
                scanned_symbols=1,
                total_matches=sum(len(items) for items in rows.values()),
                dates_with_matches=len(rows),
                downloaded_symbols=0,
                cache_hits=1,
                missing_symbols=0,
                download_failure_symbols=0,
                insufficient_symbols=0,
                universe_errors=[],
                start_date="2025-10-01",
                end_date="2025-10-03",
            )

        reports = {
            "FEEDER": report({
                "2025-10-01": [{"Symbol": "ABC", "Close": 100.0}]
            }),
            "BUY A": report({
                "2025-10-01": [{"Symbol": "ABC", "Close": 100.0}]
            }),
        }
        worker._scan = (
            lambda name, _symbols, _details, _start, _end, **_kwargs: reports[name]
        )
        worker._paper_reference_prices = lambda *_args, **_kwargs: {}

        history = worker._run_history()
        position_rule = history.paper_position_exit_rule
        self.assertEqual(position_rule["members"], ["STOP LOSS", "TAKE PROFIT"])
        self.assertEqual(position_rule["combination"], "ANY (OR)")
        self.assertEqual(position_rule["expression"]["op"], "OR")
        self.assertTrue(evaluate_position_exit_expression(
            position_rule["expression"], -8.5
        )[0])
        self.assertTrue(evaluate_position_exit_expression(
            position_rule["expression"], 16
        )[0])
        self.assertFalse(evaluate_position_exit_expression(
            position_rule["expression"], 5
        )[0])

        replay = module.replay_paper_signals(
            history.rows_by_date,
            "2025-10-01",
            "2025-10-03",
            reference_capital=100,
            position_exit_rule=position_rule,
            daily_prices={"ABC": {
                "2025-10-01": 100,
                "2025-10-02": 91.5,
                "2025-10-03": 110,
            }},
        )
        self.assertEqual(
            [row["Action"] for row in replay["Transactions"]],
            ["BUY", "SELL"],
        )
        self.assertAlmostEqual(
            replay["Transactions"][-1]["ProfitPercentAtExit"], -8.5
        )

    def test_paper_history_builds_mixed_final_set_evaluation_plan(self):
        templates = {
            name: module.default_setup()
            for name in ("FEEDER", "BUY A", "STOP LOSS", "RSI WEAK")
        }
        templates["FEEDER"]["universes"] = ["FIXED TEST SOURCE"]
        templates["STOP LOSS"]["expression"] = module.group_node("AND", [
            module.rule_node(
                module.field_operand("Loss %", "D"),
                ">=",
                module.value_operand(8.5),
            )
        ])
        templates["MIXED EXIT SET"] = {
            **module.default_setup(),
            "is_final_template": True,
            "finalset_side": "EXIT",
            "finalset_op": "AND",
            "finalset_members": ["STOP LOSS", "RSI WEAK"],
        }
        config = module.finalchartink_defaults()
        config.update({
            "recipe_name": "MIXED POSITION EXIT TEST",
            "feeder_templates": ["FEEDER"],
            "feeder_scope": "SAVED",
            "include_multirule": True,
            "buy_finalsets": ["BUY A"],
            "exit_finalsets": ["MIXED EXIT SET"],
            "as_of_mode": "historical",
            "as_of_date": "2025-10-03",
        })
        worker = module.FinalChartInkHistoryWorker(
            config,
            templates,
            {"FIXED TEST SOURCE": ("CHARTSAGE WATCHLIST", ["ABC"])},
            1,
        )

        def report(rows):
            return module.HistoryReport(
                rows_by_date=rows,
                daily_summary=[
                    {"date": day, "total": len(day_rows), "sectors": {}}
                    for day, day_rows in rows.items()
                ],
                total_symbols=1,
                scanned_symbols=1,
                total_matches=sum(len(items) for items in rows.values()),
                dates_with_matches=len(rows),
                downloaded_symbols=0,
                cache_hits=1,
                missing_symbols=0,
                download_failure_symbols=0,
                insufficient_symbols=0,
                universe_errors=[],
                start_date="2025-10-01",
                end_date="2025-10-03",
            )

        reports = {
            "FEEDER": report({
                "2025-10-01": [{"Symbol": "ABC", "Close": 100.0}]
            }),
            "BUY A": report({
                "2025-10-01": [{"Symbol": "ABC", "Close": 100.0}]
            }),
            "RSI WEAK": report({
                "2025-10-02": [{"Symbol": "ABC", "Close": 91.5}]
            }),
        }
        worker._scan = (
            lambda name, _symbols, _details, _start, _end, **_kwargs: reports[name]
        )
        worker._paper_reference_prices = lambda *_args, **_kwargs: {}

        history = worker._run_history()
        self.assertIsNone(history.paper_position_exit_rule)
        plan = history.paper_exit_evaluation_plan
        self.assertEqual(plan["combination"], "ANY (OR)")
        self.assertEqual(plan["items"], [{
            "name": "MIXED EXIT SET",
            "members": ["STOP LOSS", "RSI WEAK"],
            "required": 2,
            "combination": "AND",
        }])
        self.assertEqual(
            plan["technical_matches_by_date"]["2025-10-02"]["ABC"],
            ["RSI WEAK"],
        )

        result = module.replay_paper_signals(
            history.rows_by_date,
            "2025-10-01",
            "2025-10-03",
            reference_capital=100,
            holding_exit_rows_by_date=history.paper_exit_rows_by_date,
            exit_evaluation_plan=plan,
            daily_prices={"ABC": {
                "2025-10-01": 100,
                "2025-10-02": 91.5,
                "2025-10-03": 90,
            }},
        )
        self.assertEqual(
            [row["Action"] for row in result["Transactions"]],
            ["BUY", "SELL"],
        )
        self.assertEqual(result["Transactions"][-1]["Date"], "2025-10-02")
        self.assertEqual(
            result["Transactions"][-1]["SignalSource"],
            "Paper holding Step 3 EXIT evaluation-plan match",
        )

    def test_exit_audit_defaults_to_executed_and_filters_all_views(self):
        window = module.FinalChartInkWindow()
        try:
            window.paper_result = {
                "EXIT Evaluation Audit": [
                    {
                        "Date": "2025-01-01", "Symbol": "AAA", "Close": 100.0,
                        "Confirmations": 1, "MatchedMembers": "EXIT A",
                        "Evidence": "matched", "Executed": "Yes",
                        "ExitPnLPercent": 5.0, "Reason": "Simulated holding sold",
                    },
                    {
                        "Date": "2025-01-02", "Symbol": "BBB", "Close": 90.0,
                        "Confirmations": 1, "MatchedMembers": "EXIT A",
                        "Evidence": "matched", "Executed": "No",
                        "ExitPnLPercent": "", "Reason": "Never bought",
                    },
                    {
                        "Date": "2025-01-03", "Symbol": "CCC", "Close": 80.0,
                        "Confirmations": 1, "MatchedMembers": "EXIT A",
                        "Evidence": "matched", "Executed": "No",
                        "ExitPnLPercent": "", "Reason": "Already sold",
                    },
                ]
            }
            window._refresh_paper_exit_audit_table()
            table = window.paper_tables["EXIT Evaluation Audit"]
            self.assertEqual(window.paper_exit_audit_filter.currentData(), "EXECUTED")
            self.assertEqual(table.rowCount(), 1)
            self.assertEqual(
                window.paper_exit_audit_summary.text(),
                "EXIT signals: 3 | Executed: 1 | Not executed: 2",
            )

            window.paper_exit_audit_filter.setCurrentIndex(
                window.paper_exit_audit_filter.findData("NOT_EXECUTED")
            )
            self.app.processEvents()
            self.assertEqual(table.rowCount(), 2)
            headers = [
                table.horizontalHeaderItem(index).text()
                for index in range(table.columnCount())
            ]
            self.assertIn("Reason", headers)

            window.paper_exit_audit_filter.setCurrentIndex(
                window.paper_exit_audit_filter.findData("ALL")
            )
            self.app.processEvents()
            self.assertEqual(table.rowCount(), 3)
        finally:
            window.close()
            self.app.processEvents()

    def test_old_singular_strategy_migrates_to_one_checked_item(self):
        config = module.normalize_finalchartink_recipe({
            "recipe_name": "Old",
            "buy_finalset": "BUY A",
            "exit_finalset": "EXIT X",
        })
        self.assertEqual(config["buy_finalsets"], ["BUY A"])
        self.assertEqual(config["exit_finalsets"], ["EXIT X"])
        self.assertEqual(config["buy_finalset"], "BUY A")
        self.assertEqual(config["exit_finalset"], "EXIT X")

    def test_checkable_combos_round_trip_multiple_saved_items(self):
        window = module.FinalChartInkWindow()
        try:
            window.templates = self.templates
            window.buy_set_combo.clear()
            window.exit_set_combo.clear()
            for name in ("BUY A", "BUY B", "BUY SET"):
                window.buy_set_combo.addCheckItem(name, name)
            window.exit_set_combo.addCheckItem("EXIT X", "EXIT X")
            config = module.finalchartink_defaults()
            config.update({
                "buy_finalsets": ["BUY A", "BUY SET"],
                "exit_finalsets": ["EXIT X"],
            })
            window._apply_config(config)
            saved = window._config_from_ui()
            self.assertEqual(saved["buy_finalsets"], ["BUY A", "BUY SET"])
            self.assertEqual(saved["exit_finalsets"], ["EXIT X"])
            self.assertIn("match ANY", window.buy_set_combo.lineEdit().text())
        finally:
            window.close()
            self.app.processEvents()

    def test_mouse_clicks_can_check_more_than_one_combo_item(self):
        combo = module.CheckableComboBox(empty_text="None")
        combo.addCheckItem("Rule A", "A")
        combo.addCheckItem("Rule B", "B")
        combo.addCheckItem("Rule C", "C")
        combo.resize(360, 30)
        combo.show()
        try:
            combo.showPopup()
            self.app.processEvents()
            view = combo.view()
            for row in (0, 1):
                model_index = combo.model().index(row, 0)
                QTest.mouseClick(
                    view.viewport(),
                    module.Qt.MouseButton.LeftButton,
                    pos=view.visualRect(model_index).center(),
                )
                self.app.processEvents()
            self.assertEqual(combo.checkedData(), ["A", "B"])
            self.assertTrue(view.isVisible())

            first_index = combo.model().index(0, 0)
            QTest.mouseClick(
                view.viewport(),
                module.Qt.MouseButton.LeftButton,
                pos=view.visualRect(first_index).center(),
            )
            self.assertEqual(combo.checkedData(), ["B"])
        finally:
            combo.hidePopup()
            combo.close()
            self.app.processEvents()

    def test_quick_or_deduplicates_and_preserves_all_matching_items(self):
        worker = module.FinalChartInkCycleWorker(
            module.finalchartink_defaults(), self.templates, {}
        )
        _FakePipeline.calls = []
        with patch.object(module, "PaperPipelineWorker", _FakePipeline):
            payload = worker._run_step3_quick_or(
                ["AAA", "BBB", "CCC", "DDD"],
                ["BUY A", "BUY B", "BUY SET"],
                ["EXIT X"],
            )
        rows = {row["Symbol"]: row for row in payload["guidance_rows"]}
        self.assertEqual(set(rows), {"AAA", "BBB", "CCC", "DDD"})
        self.assertEqual(rows["AAA"]["Signal"], "BUY")
        self.assertEqual(rows["BBB"]["Signal"], "REVIEW (BUY + EXIT)")
        self.assertEqual(rows["DDD"]["Signal"], "EXIT / WEAKNESS")
        self.assertEqual(
            rows["AAA"]["Matched Step 3 Items"], "BUY A, BUY B"
        )
        self.assertEqual(rows["CCC"]["Matched Step 3 Items"], "BUY SET")
        finalset_call = next(
            call for call in _FakePipeline.calls
            if call.get("finalset_name") == "BUY SET"
        )
        self.assertEqual(finalset_call["member_names"], ["BUY A", "BUY B"])

    def test_validation_accepts_many_and_rejects_missing_item(self):
        config = module.finalchartink_defaults()
        config.update({
            "final_additional_sources": ["NIFTY 500"],
            "buy_finalsets": ["BUY A", "BUY SET"],
            "exit_finalsets": ["EXIT X"],
        })
        self.assertEqual(
            module.validate_finalchartink_cycle_config(config, self.templates), ""
        )
        config["buy_finalsets"].append("MISSING")
        self.assertIn(
            "MISSING",
            module.validate_finalchartink_cycle_config(config, self.templates),
        )

    def test_auto_names_change_when_any_checked_step3_item_changes(self):
        config = module.finalchartink_defaults()
        config.update({
            "recipe_name": "Multi",
            "final_additional_sources": ["NIFTY 500"],
            "buy_finalsets": ["BUY A", "BUY B"],
            "exit_finalsets": ["EXIT X"],
        })
        first = module.build_finalchartink_auto_watchlist_names(
            config, self.templates
        )
        changed = dict(config)
        changed["buy_finalsets"] = ["BUY A", "BUY SET"]
        second = module.build_finalchartink_auto_watchlist_names(
            changed, self.templates
        )
        self.assertNotEqual(first["watchlist_flow_id"], second["watchlist_flow_id"])
        self.assertIn("S3B-", first["final_input_watchlist"])

    def test_duplicate_strategy_is_unsaved_editable_and_independent(self):
        window = module.FinalChartInkWindow()
        try:
            source_name = "Momentum Strategy"
            source = module.finalchartink_defaults()
            source.update({
                "recipe_name": source_name,
                "auto_watchlist_naming": False,
                "multirule_watchlist": "Original Multi",
                "manual_multirule_watchlist": "Original Multi",
                "final_input_watchlist": "Original Final",
                "manual_final_input_watchlist": "Original Final",
                "interval_minutes": 27,
                "technical_adjusted": False,
            })
            existing_copy = f"{source_name} — Copy"
            window.recipe_state = {
                "version": 1,
                "active": source_name,
                "recipes": {
                    source_name: source,
                    existing_copy: {
                        **source,
                        "recipe_name": existing_copy,
                    },
                },
            }
            window.recipe_combo.blockSignals(True)
            window.recipe_combo.clear()
            window.recipe_combo.addItem("— New unsaved strategy —", "")
            window.recipe_combo.addItem(source_name, source_name)
            window.recipe_combo.addItem(existing_copy, existing_copy)
            window.recipe_combo.setCurrentIndex(1)
            window.recipe_combo.blockSignals(False)
            window._apply_config(source)
            window._update_delete_recipe_button()
            original_snapshot = dict(window.recipe_state["recipes"][source_name])

            self.assertTrue(window.duplicate_recipe_btn.isEnabled())
            window._duplicate_recipe()

            self.assertEqual(window.recipe_combo.currentData(), "")
            self.assertEqual(
                window.recipe_name_edit.text(),
                f"{source_name} — Copy 2",
            )
            self.assertEqual(window.interval_spin.value(), 27)
            self.assertFalse(window.adjusted_cb.isChecked())
            self.assertNotEqual(
                window.multirule_name_edit.text(), "Original Multi"
            )
            self.assertNotEqual(
                window.final_source_name_edit.text(), "Original Final"
            )
            self.assertFalse(window.delete_recipe_btn.isEnabled())
            self.assertFalse(window.duplicate_recipe_btn.isEnabled())
            self.assertEqual(
                window.recipe_state["recipes"][source_name], original_snapshot
            )

            window.recipe_name_edit.setText("My Modified Copy")
            self.assertEqual(window.recipe_name_edit.text(), "My Modified Copy")
            self.assertNotIn(
                "My Modified Copy", window.recipe_state["recipes"]
            )
        finally:
            window.close()
            self.app.processEvents()

    def test_latest_history_anchors_to_newest_official_bhavcopy(self):
        requested = module.NSE_CALENDAR.previous_session(module.date.today())
        official = module.NSE_CALENDAR.previous_session(
            requested - module.timedelta(days=1)
        )
        config = module.finalchartink_defaults()
        config["as_of_mode"] = "latest"
        with patch.object(
            module, "latest_official_bhavcopy_session", return_value=official
        ):
            worker = module.FinalChartInkHistoryWorker(
                config, self.templates, {}, 6
            )
        self.assertEqual(worker.config["history_requested_end"], requested.isoformat())
        self.assertEqual(worker.history_end, official.isoformat())
        self.assertEqual(worker.config["history_effective_end"], official.isoformat())
        self.assertTrue(worker.config["history_official_anchor_applied"])
        self.assertEqual(
            worker.config["history_price_date_policy"],
            "latest verified official Bhavcopy session",
        )

    def test_final_history_supports_one_two_and_three_month_periods(self):
        window = module.FinalChartInkWindow()
        try:
            choices = [
                window.history_period.itemData(index)
                for index in range(window.history_period.count())
            ]
            self.assertEqual(choices, [1, 2, 3, 6, 12, 24, 48])
        finally:
            window.close()
            self.app.processEvents()

        config = module.finalchartink_defaults()
        config.update({
            "as_of_mode": "historical",
            "as_of_date": "2026-09-30",
        })
        anchor = module.date.fromisoformat("2026-09-30")
        with patch.object(
            module.NSE_CALENDAR, "previous_session", return_value=anchor
        ):
            for months in (1, 2, 3):
                with self.subTest(months=months):
                    worker = module.FinalChartInkHistoryWorker(
                        config, self.templates, {}, months
                    )
                    expected_start = (
                        pd.Timestamp(anchor) - pd.DateOffset(months=months)
                    ).date().isoformat()
                    self.assertEqual(worker.history_start, expected_start)
                    self.assertEqual(worker.history_end, anchor.isoformat())
                    self.assertIn(
                        months, module.FINALCHARTINK_HISTORY_MONTH_CHOICES
                    )

        self.assertEqual(module.finalchartink_history_period_label(1), "1 month")
        self.assertEqual(module.finalchartink_history_period_label(2), "2 months")
        self.assertEqual(module.finalchartink_history_period_label(3), "3 months")

    def test_every_final_xlsx_path_carries_current_strategy_descriptions(self):
        config = module.finalchartink_defaults()
        config.update({
            "recipe_name": "XLSX Messaging",
            "feeder_templates": ["BUY A"],
            "feeder_scope": "SAVED",
            # Deliberately check a Step 2A source so the export must explain
            # that this choice is ignored by SAVED mode.
            "feeder_sources": ["NIFTY 500"],
            "include_multirule": True,
            "final_additional_sources": [],
            "buy_finalsets": ["BUY A", "BUY B"],
            "exit_finalsets": ["EXIT X"],
            "history_months": 1,
        })
        selected_rules = module.build_finalchartink_selected_rule_rows(
            config, self.templates
        )

        def sheet_text(sheet):
            return "\n".join(
                str(value)
                for row in sheet.iter_rows(values_only=True)
                for value in row
                if value is not None
            )

        expected_messages = (
            "Use each Step 1 template's saved universe only",
            module.FINALCHARTINK_STEP2A_SAVED_SELECTION_MESSAGE,
            module.FINALCHARTINK_STEP2B_EMPTY_SELECTION_MESSAGE,
            "BUY A",
            "BUY B",
            "EXIT X",
        )

        live_path = ROOT / "build" / "_test_final_xlsx_descriptions.xlsx"
        history_path = ROOT / "build" / "_test_final_history_descriptions.xlsx"
        live_path.parent.mkdir(parents=True, exist_ok=True)
        try:
            module.export_finalchartink_workbook(
                live_path,
                {
                    "config": config,
                    "guidance_rows": [],
                    "selected_rule_rows": selected_rules,
                    "timestamped_watchlists": [],
                    "feeder_results": [],
                    "feeder_match_rows": [],
                },
            )
            live_workbook = load_workbook(live_path, data_only=True)
            try:
                self.assertIn("Run Summary", live_workbook.sheetnames)
                self.assertIn("How Strategy Works", live_workbook.sheetnames)
                live_text = "\n".join((
                    sheet_text(live_workbook["Run Summary"]),
                    sheet_text(live_workbook["How Strategy Works"]),
                ))
                for message in expected_messages:
                    self.assertIn(message, live_text)
            finally:
                live_workbook.close()

            report = module.HistoryReport(
                rows_by_date={},
                daily_summary=[],
                total_symbols=0,
                scanned_symbols=0,
                total_matches=0,
                dates_with_matches=0,
                downloaded_symbols=0,
                cache_hits=0,
                missing_symbols=0,
                download_failure_symbols=0,
                insufficient_symbols=0,
                universe_errors=[],
                start_date="2026-09-01",
                end_date="2026-09-30",
            )
            module.export_finalchartink_history_workbook(
                str(history_path), report, config, self.templates
            )
            history_workbook = load_workbook(history_path, data_only=True)
            try:
                self.assertIn("Strategy Flow & Rules", history_workbook.sheetnames)
                self.assertIn("How Strategy Works", history_workbook.sheetnames)
                history_text = "\n".join((
                    sheet_text(history_workbook["Strategy Flow & Rules"]),
                    sheet_text(history_workbook["How Strategy Works"]),
                ))
                for message in expected_messages:
                    self.assertIn(message, history_text)
                self.assertIn("1 month", history_text)
            finally:
                history_workbook.close()
        finally:
            live_path.unlink(missing_ok=True)
            history_path.unlink(missing_ok=True)

    def test_historical_as_of_keeps_its_selected_session_policy(self):
        config = module.finalchartink_defaults()
        config.update({
            "as_of_mode": "historical",
            "as_of_date": "2026-09-20",
        })
        expected = module.NSE_CALENDAR.previous_session(
            module.date.fromisoformat("2026-09-20")
        )
        with patch.object(
            module,
            "latest_official_bhavcopy_session",
            side_effect=AssertionError("Historical mode must not consult latest"),
        ):
            worker = module.FinalChartInkHistoryWorker(
                config, self.templates, {}, 6
            )
        self.assertEqual(worker.history_end, expected.isoformat())
        self.assertFalse(worker.config["history_official_anchor_applied"])
        self.assertEqual(
            worker.config["history_price_date_policy"],
            "selected historical session",
        )

    def test_component_history_scan_uses_effective_date_exactly(self):
        config = module.finalchartink_defaults()
        config["as_of_mode"] = "latest"
        official = module.date.fromisoformat("2026-09-23")
        with patch.object(
            module, "latest_official_bhavcopy_session", return_value=official
        ):
            worker = module.FinalChartInkHistoryWorker(
                config, self.templates, {}, 6
            )
        captured = {}

        class FakeHistoryScanner:
            def __init__(self, setup, **kwargs):
                captured["setup"] = setup

            def scan(self, progress):
                return "report"

        with patch.object(module, "HistoryScannerWorker", FakeHistoryScanner):
            report = worker._scan("BUY A", ["AAA"], {}, 0, 100)
        self.assertEqual(report, "report")
        self.assertEqual(captured["setup"]["as_of_mode"], "historical")
        self.assertEqual(captured["setup"]["as_of_date"], "2026-09-23")

    def test_history_frame_pool_loads_maximum_range_once_and_returns_exact_slices(self):
        dates = pd.date_range("2025-01-01", "2025-01-10", freq="D")
        full = pd.DataFrame({
            "Open": range(10), "High": range(1, 11), "Low": range(10),
            "Close": range(1, 11), "Volume": [100] * 10,
        }, index=dates)

        class FakeCache:
            def __init__(self):
                self.loads = []

            def load(self, symbol, start, end, adjusted):
                self.loads.append((symbol, pd.Timestamp(start), pd.Timestamp(end)))
                return full.loc[(full.index >= start) & (full.index <= end)].copy()

            def is_covered(self, symbol, start, end, adjusted):
                return True

        shared = {}
        cache = FakeCache()
        first = module.HistoryScannerWorker(
            module.default_setup(),
            cache_path=ROOT / "price_cache.db",
            shared_history_frames=shared,
            shared_history_start="2025-01-01",
        )
        first_frame, first_covered, first_reused = first._load_history_frame(
            cache, "ABC", pd.Timestamp("2025-01-05"),
            pd.Timestamp("2025-01-10"), True,
        )
        second = module.HistoryScannerWorker(
            module.default_setup(),
            cache_path=ROOT / "price_cache.db",
            shared_history_frames=shared,
            shared_history_start="2025-01-01",
        )
        second_frame, second_covered, second_reused = second._load_history_frame(
            cache, "ABC", pd.Timestamp("2025-01-08"),
            pd.Timestamp("2025-01-10"), True,
        )
        self.assertEqual(len(cache.loads), 1)
        self.assertTrue(first_covered)
        self.assertFalse(first_reused)
        self.assertEqual(first_frame.index.min(), pd.Timestamp("2025-01-05"))
        self.assertTrue(second_covered)
        self.assertTrue(second_reused)
        self.assertEqual(second_frame.index.min(), pd.Timestamp("2025-01-08"))

    def test_history_frame_pool_falls_back_to_original_exact_range(self):
        dates = pd.date_range("2025-01-01", "2025-01-10", freq="D")
        full = pd.DataFrame({
            "Open": range(10), "High": range(1, 11), "Low": range(10),
            "Close": range(1, 11), "Volume": [100] * 10,
        }, index=dates)

        class FakeCache:
            def __init__(self):
                self.loads = []

            def load(self, symbol, start, end, adjusted):
                self.loads.append(pd.Timestamp(start))
                return full.loc[(full.index >= start) & (full.index <= end)].copy()

            def is_covered(self, symbol, start, end, adjusted):
                return pd.Timestamp(start) >= pd.Timestamp("2025-01-05")

        shared = {}
        cache = FakeCache()
        worker = module.HistoryScannerWorker(
            module.default_setup(),
            cache_path=ROOT / "price_cache.db",
            shared_history_frames=shared,
            shared_history_start="2025-01-01",
        )
        frame, covered, reused = worker._load_history_frame(
            cache, "ABC", pd.Timestamp("2025-01-05"),
            pd.Timestamp("2025-01-10"), True,
        )
        self.assertEqual(cache.loads, [
            pd.Timestamp("2025-01-01"), pd.Timestamp("2025-01-05")
        ])
        self.assertTrue(covered)
        self.assertFalse(reused)
        self.assertFalse(shared)
        self.assertEqual(frame.index.min(), pd.Timestamp("2025-01-05"))

    def test_history_ranking_filters_dated_membership_before_top_n(self):
        setup = module.default_setup()
        setup.update({
            "as_of_mode": "historical",
            "as_of_date": "2026-09-30",
            "history_months": 1,
            "expression": module.group_node("AND", [
                module.rule_node(
                    module.field_operand("Close", "D"),
                    ">",
                    module.value_operand(0),
                )
            ]),
            "ranking_enabled": True,
            "ranking_direction": "DESC",
            "ranking_rsi_periods": [2, 3],
            "ranking_top_n": 1,
        })
        dates = pd.date_range("2025-05-01", "2026-09-30", freq="B")
        frames = {
            symbol: pd.DataFrame({
                "Open": [100.0] * len(dates),
                "High": [101.0] * len(dates),
                "Low": [99.0] * len(dates),
                "Close": [100.0] * len(dates),
                "Volume": [1000.0] * len(dates),
            }, index=dates)
            for symbol in ("CURRENT", "FORMER")
        }

        class FakeCache:
            COLUMNS = ["Open", "High", "Low", "Close", "Volume"]

            def __init__(self, *_args, **_kwargs):
                pass

            def load(self, symbol, start, end, _adjusted):
                frame = frames[symbol]
                return frame.loc[(frame.index >= start) & (frame.index <= end)].copy()

            def is_covered(self, *_args, **_kwargs):
                return True

            def candle_source(self, *_args, **_kwargs):
                return "Test"

        scored_symbols = []

        def fake_ranking_values(row, _setup, _value_for_period):
            symbol = str(row["Symbol"])
            scored_symbols.append(symbol)
            row[module.COMPOSITE_RSI_MOMENTUM_COLUMN] = (
                100.0 if symbol == "FORMER" else 50.0
            )

        worker = module.HistoryScannerWorker(
            setup,
            symbols=["CURRENT", "FORMER"],
            symbol_details={},
            cache_path=ROOT / "price_cache.db",
            dated_symbol_filter=lambda symbol, _stamp: symbol == "CURRENT",
        )
        with (
            patch.object(module, "PriceCache", FakeCache),
            patch.object(
                module,
                "add_composite_rsi_ranking_values",
                side_effect=fake_ranking_values,
            ),
        ):
            report = worker.scan()

        self.assertTrue(report.rows_by_date)
        self.assertEqual(
            {
                row["Symbol"]
                for rows in report.rows_by_date.values()
                for row in rows
            },
            {"CURRENT"},
        )
        self.assertNotIn("FORMER", scored_symbols)

    def test_compact_telegram_lists_close_and_daily_change(self):
        payload = {
            "config": {"recipe_name": "Test & Strategy"},
            "guidance_rows": [
                {
                    "Symbol": "INDGN",
                    "Signal": "BUY",
                    "Close": 591.75,
                    module.PERCENT_CHANGE_COLUMN: 1.843,
                },
                {
                    "Symbol": "M&M",
                    "Signal": "EXIT / WEAKNESS",
                    "Close": 2048.5,
                    module.PERCENT_CHANGE_COLUMN: -0.426,
                },
                {
                    "Symbol": "DUAL",
                    "Signal": "REVIEW",
                    "Close": None,
                    module.PERCENT_CHANGE_COLUMN: None,
                },
                {
                    "Symbol": "DUAL",
                    "Signal": "BUY",
                    "Close": 100.0,
                    module.PERCENT_CHANGE_COLUMN: 5.0,
                },
            ],
        }
        generated = module.datetime(2026, 9, 26, 11, 33, 19)
        with patch.object(module, "get_ist_now", return_value=generated):
            messages = module.build_finalchartink_compact_telegram_messages(payload)
        text = "\n".join(messages)

        self.assertIn("Strategy: Test &amp; Strategy", text)
        self.assertIn(
            "<b>1. INDGN</b> — Close: ₹591.75 | Change: +1.84%", text
        )
        self.assertIn(
            "<b>1. M&amp;M</b> — Close: ₹2,048.5 | Change: -0.43%", text
        )
        self.assertIn(
            "<b>1. DUAL</b> — Close: n/a | Change: n/a", text
        )
        self.assertEqual(text.count("<b>1. DUAL</b>"), 1)
        self.assertNotIn("<b>Close:", text)


if __name__ == "__main__":
    unittest.main()
